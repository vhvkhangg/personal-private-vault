"""Bounded local development workload; no held-out, publication or GPU acceptance claim."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import socket
import sys
import tempfile
import threading
import time
from dataclasses import asdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "rag-engine/src"))

import docx  # noqa: E402
import onnxruntime as ort  # noqa: E402
import openpyxl  # noqa: E402
import pypdf  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402

from ppv_rag.chunking.chunker import StructureFirstChunker  # noqa: E402
from ppv_rag.embeddings.model import MultilingualEmbeddingPipeline  # noqa: E402
from ppv_rag.limits import parsing_scope  # noqa: E402
from ppv_rag.ocr.engine import AdaptiveOCREngine  # noqa: E402
from ppv_rag.parsers.docx import parse_docx  # noqa: E402
from ppv_rag.parsers.excel import parse_xlsx  # noqa: E402
from ppv_rag.parsers.markdown import parse_markdown  # noqa: E402
from ppv_rag.parsers.pdf import parse_pdf  # noqa: E402
from ppv_rag.telemetry import (  # noqa: E402
    get_gpu_vram_bytes,
    get_host_ram_headroom_bytes,
    get_host_ram_total_bytes,
    get_system_ram_bytes,
    gpu_measurement_available,
)


def saved(document: object) -> bytes:
    output = io.BytesIO()
    document.save(output)  # type: ignore[attr-defined]
    return output.getvalue()


def scanned_pdf(pages: int) -> bytes:
    image = Image.new("RGB", (1400, 400), "white")
    ImageDraw.Draw(image).text(
        (50, 100),
        "Resource workload contains 45 books.",
        font=ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 42),
        fill="black",
    )
    output = io.BytesIO()
    image.save(output, "PDF", resolution=144)
    page = pypdf.PdfReader(io.BytesIO(output.getvalue())).pages[0]
    writer = pypdf.PdfWriter()
    for _ in range(pages):
        writer.add_page(page)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--termination-workload", action="store_true")
    args = parser.parse_args()

    # Guard tested Python application paths. Native OS egress confinement is not claimed.
    def deny(*_args: object, **_kwargs: object) -> None:
        raise RuntimeError("NETWORK_CALL_DENIED")

    socket.socket = deny  # type: ignore[assignment]
    if args.termination_workload:
        ocr = AdaptiveOCREngine()
        run_pass = ocr._run_pass
        entered = False

        def measured_pass(*args: object, **kwargs: object) -> object:
            nonlocal entered
            if not entered:
                print(
                    json.dumps({"pid": os.getpid(), "phase": "REAL_OCR_NATIVE_PASS_ENTERED"}),
                    flush=True,
                )
                entered = True
            return run_pass(*args, **kwargs)  # type: ignore[arg-type]

        ocr._run_pass = measured_pass  # type: ignore[method-assign,assignment]
        with parsing_scope():
            parse_pdf(scanned_pdf(100), "termination-development", 1, ocr_engine=ocr)
        raise RuntimeError("Termination diagnostic unexpectedly completed before watchdog")

    started, cpu = time.perf_counter(), time.process_time()
    samples: list[int] = []
    device_samples: list[int] = []
    pipeline_disk_samples: list[int] = []
    phase = "fixture_generation"
    stop = threading.Event()
    with tempfile.TemporaryDirectory(prefix="ppv-s2-development-") as scratch:
        previous = tempfile.tempdir
        tempfile.tempdir = scratch
        old_env = {name: os.environ.get(name) for name in ("TMP", "TEMP", "TMPDIR")}
        for name in old_env:
            os.environ[name] = scratch

        def sample() -> None:
            while not stop.is_set():
                samples.append(
                    sum(p.stat().st_size for p in Path(scratch).rglob("*") if p.is_file())
                )
                if phase == "pipeline":
                    pipeline_disk_samples.append(samples[-1])
                if gpu_measurement_available():
                    device_samples.append(get_gpu_vram_bytes())
                stop.wait(0.02)

        monitor = threading.Thread(target=sample)
        monitor.start()
        try:
            ocr = AdaptiveOCREngine()
            scan_bytes = scanned_pdf(3)
            word = docx.Document()
            for index in range(200):
                word.add_paragraph(f"Development paragraph {index}: dữ liệu; 中文.")
            workbook = openpyxl.Workbook()
            for index in range(500):
                workbook.active.append([f"Row {index}", index, "=SUM(1,2)"])
            word_bytes, workbook_bytes = saved(word), saved(workbook)
            phase = "pipeline"
            with parsing_scope(seconds=120):
                documents = [
                    parse_docx(word_bytes, "word", 1),
                    parse_xlsx(workbook_bytes, "sheet", 1),
                    parse_markdown(
                        "---\nprivate: true\n---\n" + "text\n" * 200, "md", 1, whole_file=True
                    ),
                    parse_pdf(scan_bytes, "scan", 1, ocr_engine=ocr),
                ]
                model = MultilingualEmbeddingPipeline(force_cpu=True)
                chunks = [
                    chunk
                    for document in documents
                    for chunk in StructureFirstChunker(token_count=model.token_count).chunk(
                        document
                    )
                ]
                vectors = model.embed_batch([chunk.content for chunk in chunks])
            print(
                json.dumps(
                    {
                        "scope": "EXPOSED_DEVELOPMENT_ONLY; not held-out acceptance",
                        "pid": os.getpid(),
                        "documents": len(documents),
                        "elements": sum(len(d.elements) for d in documents),
                        "chunks": len(chunks),
                        "vectors": len(vectors),
                        "wall_seconds": time.perf_counter() - started,
                        "process_cpu_seconds": time.process_time() - cpu,
                        "process_peak_working_set_bytes": get_system_ram_bytes() or None,
                        "host_headroom_bytes_at_finish": get_host_ram_headroom_bytes() or None,
                        "host_total_ram_bytes": get_host_ram_total_bytes() or None,
                        "device_wide_vram_max_sample_bytes": max(device_samples)
                        if device_samples
                        else None,
                        "sampling_period_seconds": 0.02,
                        "sample_count": len(samples),
                        "scratch_disk_max_sample_bytes": max(samples) if samples else None,
                        "pipeline_scratch_disk_max_sample_bytes": (
                            max(pipeline_disk_samples) if pipeline_disk_samples else None
                        ),
                        "disk_scope_includes_fixture_generation": True,
                        "scratch_disk_final_bytes": sum(
                            p.stat().st_size for p in Path(scratch).rglob("*") if p.is_file()
                        ),
                        "disk_scope": (
                            "dedicated TEMP/TMP/TMPDIR; sampled, not kernel write accounting"
                        ),
                        "cuda": "UNAVAILABLE"
                        if "CUDAExecutionProvider" not in ort.get_available_providers()
                        else "AVAILABLE_NOT_RUN",
                        "providers": ort.get_available_providers(),
                        "pipeline": asdict(model.metadata()),
                        "fixture_sha256": hashlib.sha256(scan_bytes).hexdigest(),
                    },
                    indent=2,
                ),
                flush=True,
            )
        finally:
            stop.set()
            monitor.join(5)
            tempfile.tempdir = previous
            for name, value in old_env.items():
                if value is None:
                    os.environ.pop(name, None)
                else:
                    os.environ[name] = value
    print(json.dumps({"scratch_directory_removed": not Path(scratch).exists()}), flush=True)


if __name__ == "__main__":
    main()
