"""Create a new, immutable, synthetic held-out/adjudication package, without running OCR.

Labels come from authoring instructions and original structures, never parser/OCR output.
Human signatures are deliberately absent. Do not regenerate or tune against this dataset.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
from collections import Counter
from pathlib import Path
from typing import Any

import docx  # type: ignore[import-untyped]
import openpyxl
import pypdf
import xlwt  # type: ignore[import-untyped]
from openpyxl.comments import Comment
from PIL import Image, ImageDraw, ImageFont
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

ROOT = Path(__file__).resolve().parents[2]


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save_json(path: Path, value: object) -> None:
    # Fixture hashes must survive the repository's eol=lf Git normalization on Windows.
    path.write_bytes((json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))


def native_pdf(text: str) -> bytes:
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(612, 792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Courier"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 72 700 Td ({text}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(content)
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


def build(output: Path) -> None:
    if output.exists():
        raise ValueError("Immutable dataset destination already exists; choose a NEW revision")
    output.mkdir(parents=True)
    docs: list[dict[str, Any]] = []
    font_hashes: dict[str, str] = {}

    def add(
        name: str,
        family: str,
        fmt: str,
        scripts: list[str],
        gold: list[dict[str, Any]],
        *,
        pages: int = 0,
        tables: int = 0,
        cells: int = 0,
        dependencies: list[str] | None = None,
        expected: str = "EXACT_MATERIAL_AND_LOCATORS_OR_FAIL",
    ) -> None:
        docs.append(
            dict(
                file=name,
                sha256=sha(output / name),
                source_family=family,
                closure_id=family,
                split="held_out",
                format=fmt,
                scripts=scripts,
                source_version=1,
                whole_file_consent_fixture=True,
                source_key=f"s2-new-{name}",
                original_digest=sha(output / name),
                expected=expected,
                dependencies=dependencies or [],
                denominators=dict(
                    documents=1, pages=pages, blocks=len(gold), tables=tables, cells=cells
                ),
                gold=gold,
            )
        )

    texts = {
        "en": ["Courier record 708", "Recipient: Helen Brooks.", "27 parcels at 08:45."],
        "vi": ["Phiếu chuyển hàng 708", "Người nhận: Đỗ Thị Hồng.", "Số kiện 27, giờ 08:45."],
        "zh-Hans": ["配送记录708", "收件人：陈晓东。", "件数27，时间08:45。"],
        "zh-Hant": ["配送紀錄708", "收件人：陳曉東。", "件數27，時間08:45。"],
    }
    for script, lines in texts.items():
        font_path = Path(
            "C:/Windows/Fonts/msyh.ttc" if script.startswith("zh") else "C:/Windows/Fonts/arial.ttf"
        )
        font_hashes[font_path.name] = sha(font_path)
        font = ImageFont.truetype(str(font_path), 36)
        image = Image.new("RGB", (1600, 900), "white")
        draw = ImageDraw.Draw(image)
        gold = []
        for index, text in enumerate(lines):
            point = (80, 60 + index * 75)
            draw.text(point, text, font=font, fill="black")
            gold.append(
                dict(
                    text=text,
                    reading_order=index,
                    page=1,
                    coordinate_system="PDF_TOP_LEFT",
                    ink_bbox=list(draw.textbbox(point, text, font=font)),
                    provenance="source font renderer ink box; human review pending",
                )
            )
        # Original ruled grid: merged first row, followed by two two-column rows.
        for y in (350, 450, 550, 650):
            draw.line((80, y, 1440, y), fill="black", width=3)
        for x in (80, 1440):
            draw.line((x, 350, x, 650), fill="black", width=3)
        draw.line((760, 450, 760, 650), fill="black", width=3)
        for index, (text, row, col, x0, x1) in enumerate(
            [
                ("Dispatch", 0, 0, 80, 1440),
                ("Code", 1, 0, 80, 760),
                ("Count", 1, 1, 760, 1440),
                ("708", 2, 0, 80, 760),
                ("27", 2, 1, 760, 1440),
            ]
        ):
            y0 = 350 + row * 100
            draw.text((x0 + 24, y0 + 24), text, font=font, fill="black")
            gold.append(
                dict(
                    text=text,
                    reading_order=3 + index,
                    page=1,
                    coordinate_system="PDF_TOP_LEFT",
                    table=0,
                    row=row,
                    column=col,
                    column_span=2 if row == 0 else 1,
                    bbox=[x0, y0, x1, y0 + 100],
                    ink_bbox=list(draw.textbbox((x0 + 24, y0 + 24), text, font=font)),
                )
            )
        name = f"courier-{script}.pdf"
        image.save(output / name, "PDF", resolution=72)
        add(name, "new-courier-record", "pdf-scan", [script], gold, pages=1, tables=1, cells=5)

    source_text = "Workshop session begins at 16:20."
    (output / "workshop-native.pdf").write_bytes(native_pdf(source_text))
    native_gold = [
        dict(
            text=source_text,
            page=1,
            coordinate_system="PDF_TOP_LEFT",
            bbox=[72, 84.452, 72 + len(source_text) * 7.2, 93.884],
            font="Courier",
            font_size=12,
            source_baseline=[72, 700],
            provenance="authored PDF content stream and Courier metrics; human review pending",
        )
    ]
    add("workshop-native.pdf", "new-workshop-notice", "pdf-native", ["en"], native_gold, pages=1)
    writer = pypdf.PdfWriter()
    writer.append(io.BytesIO((output / "workshop-native.pdf").read_bytes()))
    writer.append(io.BytesIO((output / "courier-vi.pdf").read_bytes()))
    with (output / "courier-mixed.pdf").open("wb") as stream:
        writer.write(stream)
    # This compound source belongs to a closure containing BOTH families, all in held_out.
    add(
        "courier-mixed.pdf",
        "new-courier-workshop-closure",
        "pdf-mixed",
        ["en", "vi"],
        native_gold + [{**g, "page": 2} for g in docs[1]["gold"]],
        pages=2,
        tables=1,
        cells=5,
        dependencies=["workshop-native.pdf", "courier-vi.pdf"],
    )

    text = "Lịch sửa máy ngày 24: đổi dây curoa, kiểm tra 63 vít.\nGhi chú: Đặng Thu Hà.\n"
    (output / "maintenance.txt").write_bytes(text.encode("utf-8"))
    add(
        "maintenance.txt",
        "new-maintenance-notebook",
        "txt",
        ["vi"],
        [dict(text=text, char_start=0, char_end=len(text), line_start=1, line_end=2)],
    )
    child = "# 工具\n扳手編號92，保留隱藏備註。 ^tool92\n"
    parent = (
        "---\nprivate: true\ncustom: 安全檢查\n---\n# 維修筆記\n"
        "%% 核對螺絲63枚 %%\n![[tools#^tool92]]\n"
    )
    (output / "tools.md").write_bytes(child.encode("utf-8"))
    (output / "maintenance.md").write_bytes(parent.encode("utf-8"))
    add(
        "tools.md",
        "new-maintenance-notebook",
        "md/Obsidian",
        ["zh-Hant"],
        [dict(text=child, original_line_range=[1, 2], block_id="tool92")],
    )
    add(
        "maintenance.md",
        "new-maintenance-notebook",
        "md/Obsidian",
        ["zh-Hant"],
        [
            dict(
                text=parent,
                original_line_range=[1, len(parent.splitlines())],
                includes_frontmatter=True,
                includes_comment=True,
            ),
            dict(
                text="扳手編號92，保留隱藏備註。",
                child_file="tools.md",
                child_source_key="s2-new-tools.md",
                child_source_version=1,
                child_original_digest=sha(output / "tools.md"),
                child_original_line=2,
                block_id="tool92",
            ),
        ],
        dependencies=["tools.md"],
    )

    word = docx.Document()
    word.core_properties.title = "Maintenance instruction 92"
    paragraph = word.add_paragraph("Kiểm tra nhiệt độ 37 độ.")
    hidden = paragraph.add_run(" Chỉ dẫn ẩn: thay vít 63.")
    hidden.font.hidden = True
    word.add_comment(paragraph.runs[0], "Đã kiểm tra ngày 24.", author="Synthetic label author")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Tool 92", "Count 63"
    word.save(output / "maintenance.docx")
    add(
        "maintenance.docx",
        "new-maintenance-notebook",
        "docx",
        ["vi", "en"],
        [
            dict(
                text="Maintenance instruction 92",
                container_part="core_properties",
                property_key="title",
            ),
            dict(text=paragraph.text, paragraph_index=0, contains_hidden=True),
            dict(
                text="Đã kiểm tra ngày 24.",
                comment_id="0",
                anchored_paragraph=0,
                anchor_char_range=[0, len(paragraph.runs[0].text)],
            ),
            dict(text="Tool 92", table=0, row=0, column=0),
            dict(text="Count 63", table=0, row=0, column=1),
        ],
        tables=1,
        cells=2,
    )

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Inspection"
    sheet["A1"] = "Inspection 92"
    sheet.merge_cells("A1:C1")
    sheet["A2"], sheet["C2"] = "螺絲", 63
    sheet["C2"].comment = Comment("核對數量63", "Synthetic")
    sheet.column_dimensions.group("B", "D", hidden=True)
    sheet["A3"] = "=SUM(1,2)"
    workbook.save(output / "inspection.xlsx")
    add(
        "inspection.xlsx",
        "new-inspection-register",
        "xlsx",
        ["en", "zh-Hant"],
        [
            dict(
                text="Inspection 92", worksheet="Inspection", a1_range="A1:C1", merged_anchor="A1"
            ),
            dict(text="螺絲", worksheet="Inspection", a1_range="A2"),
            dict(text="63", worksheet="Inspection", a1_range="C2", hidden=True),
            dict(text="核對數量63", worksheet="Inspection", a1_range="C2", comment=True),
            dict(
                text="=SUM(1,2)",
                worksheet="Inspection",
                a1_range="A3",
                formula_state="UNCACHED_INERT_TEXT",
            ),
        ],
        tables=1,
        cells=5,
    )
    legacy = xlwt.Workbook()
    sheet2 = legacy.add_sheet("登记")
    sheet2.write_merge(0, 0, 0, 1, "零件登记")
    sheet2.write(1, 0, "扳手92")
    sheet2.write(1, 1, 63)
    sheet2.col(1).hidden = True
    legacy.save(str(output / "inspection.xls"))
    add(
        "inspection.xls",
        "new-inspection-register",
        "xls",
        ["zh-Hans"],
        [
            dict(text="零件登记", worksheet="登记", a1_range="A1:B1", merged_anchor="A1"),
            dict(text="扳手92", worksheet="登记", a1_range="A2"),
            dict(text="63", worksheet="登记", a1_range="B2", hidden=True),
        ],
        tables=1,
        cells=3,
    )

    for name, fmt in (("broken.docx", "docx"), ("broken.pdf", "pdf-scan")):
        (output / name).write_bytes(b"Not a valid document; must fail closed.\x00")
        add(name, "new-corrupt-negative", fmt, ["en"], [], expected="FILE_FAILURE_NO_OUTPUT")
    save_json(
        output / "missing-locator.json",
        {
            "text": "Seeded material",
            "locator": {"page": 1},
            "expected": "REJECT_MISSING_SOURCE_REVISION_AND_BBOX",
        },
    )
    add(
        "missing-locator.json",
        "new-locator-negative",
        "locator-negative",
        ["en"],
        [],
        expected="MANDATORY_LOCATOR_REJECTION",
    )

    # All linked sources must be held_out, including the mixed-document union.
    names = {d["file"] for d in docs}
    assert all(set(d["dependencies"]) <= names for d in docs)
    original = json.loads(
        (ROOT / "rag-engine/tests/fixtures/phase18-ocr-v2/manifest.json").read_text(
            encoding="utf-8"
        )
    )
    old_hashes = {d["sha256"] for d in original["items"]}
    assert not old_hashes.intersection(d["sha256"] for d in docs)
    totals = {
        key: sum(d["denominators"][key] for d in docs)
        for key in ("documents", "pages", "blocks", "tables", "cells")
    }
    by_format = {}
    by_script = {}
    for dimension, target in (("format", by_format), ("scripts", by_script)):
        for key in sorted(
            {
                x
                for d in docs
                for x in (d[dimension] if isinstance(d[dimension], list) else [d[dimension]])
            }
        ):
            members = [
                d
                for d in docs
                if key in (d[dimension] if isinstance(d[dimension], list) else [d[dimension]])
            ]
            target[key] = {
                metric: sum(d["denominators"][metric] for d in members) for metric in totals
            }
    save_json(
        output / "manifest.json",
        dict(
            dataset=output.name,
            synthetic=True,
            reference_origin=(
                "authored source strings, original OOXML/BIFF/PDF objects, "
                "font-renderer ink boxes; NEVER OCR/parser output"
            ),
            human_adjudication="PENDING",
            pipeline_selection=(
                "NONE; candidate comparison incomplete; do not evaluate held-out yet"
            ),
            held_out_results="NOT_RUN",
            fonts=font_hashes,
            documents=docs,
            denominators=totals,
            by_format=by_format,
            by_script=by_script,
            group_count=len({d["source_family"] for d in docs}),
            cross_split_original_hash_intersection=[],
            family_assignment=(
                "all new translations/scans/exports and complete child/mixed closures held_out; "
                "old shelf/invoice and public upstream samples development only"
            ),
            near_duplicate_review=(
                "PENDING HUMAN REVIEW; no semantic independence proof from hashes alone"
            ),
        ),
    )
    (output / "manifest.sha256").write_bytes((sha(output / "manifest.json") + "\n").encode("ascii"))
    save_json(
        output / "adjudication.json",
        dict(
            status="PENDING",
            reviewer=None,
            signed_at=None,
            dataset_sha256=sha(output / "manifest.json"),
            required_checks=[
                (
                    "Read source files and gold side-by-side; correct labels only in a NEW "
                    "dataset revision before evaluation"
                ),
                (
                    "Check vi diacritics/digits and zh-Hans versus zh-Hant; assess English "
                    "whitespace without silently normalizing facts"
                ),
                "Check PDF ink/ruled/merged boxes, reading order; verify native Courier metrics",
                "Check hidden/comment/formula/child original structure and closure IDs",
                (
                    "Independently assess source-family/semantic near-duplicates "
                    "and label denominators"
                ),
                (
                    "Freeze accepted pipeline/config/artifact hashes, then independent held-out "
                    "execution; adjudicate every discrepancy and second-check "
                    "acceptance-critical ambiguities"
                ),
            ],
            documents=[
                dict(
                    file=d["file"],
                    label_status="PENDING_HUMAN_REVIEW",
                    material_error_decision=None,
                    locator_decision=None,
                    reviewer=None,
                )
                for d in docs
            ],
        ),
    )
    print(
        json.dumps(
            dict(
                output=str(output),
                totals=totals,
                formats=dict(Counter(d["format"] for d in docs)),
                manifest_sha256=sha(output / "manifest.json"),
            )
        )
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    build(parser.parse_args().output)
