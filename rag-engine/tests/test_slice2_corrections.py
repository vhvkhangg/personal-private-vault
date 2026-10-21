"""Development regressions for S2-01–08; these exposed fixtures are not held-out quality."""

import base64
import io
import math
import threading
import time
import zipfile

import docx
import numpy as np
import openpyxl
import pytest
import xlwt
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw, ImageFont

from ppv_rag.chunking.chunker import StructureFirstChunker
from ppv_rag.contracts import DocumentBatchStageRequest, DocumentRootTask, document_input_digest
from ppv_rag.embeddings.model import MultilingualEmbeddingPipeline
from ppv_rag.limits import ParsingCancelled, checked_zip, parsing_scope
from ppv_rag.ocr.engine import OCRResult, OCRTextBlock
from ppv_rag.parsers.docx import parse_docx
from ppv_rag.parsers.excel import parse_xls, parse_xlsx
from ppv_rag.parsers.markdown import ApprovedDependency, parse_markdown
from ppv_rag.parsers.models import DocumentElement, ElementKind
from ppv_rag.parsers.pdf import parse_pdf
from ppv_rag.parsers.text import parse_text


def saved(document):
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_producer_components_import_in_fresh_process_without_order_dependency():
    import os
    import subprocess
    import sys
    from pathlib import Path

    result = subprocess.run(
        [sys.executable, "-c", "import ppv_rag.ocr.engine; import ppv_rag.embeddings.model"],
        env={**os.environ, "PYTHONPATH": str(Path("rag-engine/src").resolve())},
        capture_output=True,
        text=True,
        timeout=15,
        check=False,
    )
    assert result.returncode == 0, result.stderr


WHOLE_FILE = (
    '---\nprivate: true\nnsfw: true\ntitle: "Original title"\ncustom: value\n---\n'
    "%% Approved comment %%\n# Body\n<!-- Whole comment -->\n"
    "```html\n<!-- Literal -->\n![[inert-code-embed]]\n```\n"
)


def test_whole_file_markdown_retains_original_frontmatter_comments_and_code():
    parsed = parse_markdown(WHOLE_FILE, "parent", 5, whole_file=True)
    assert parsed.elements[0].text == WHOLE_FILE.split("%%")[0]
    assert "".join(e.text for e in parsed.elements) == WHOLE_FILE
    assert all(e.approved_field == "content" for e in parsed.elements)
    assert "Whole comment" in "\n".join(e.text for e in parsed.elements)
    code = next(e for e in parsed.elements if e.kind == ElementKind.CODE_BLOCK)
    assert code.text == "```html\n<!-- Literal -->\n![[inert-code-embed]]\n```\n"
    assert code.locator["line_start"] == 10
    # The separate field-filtered policy retains restrictive denial.
    with pytest.raises(ValueError, match="Restrictive"):
        parse_markdown(WHOLE_FILE, "domain-style-fields", 1, whole_file=False)


@pytest.mark.parametrize(
    "content", ["---\nx: 1\nx: 2\n---\nbody", "---\nx: 1\nbody", "---\na: &a [x]\nb: *a\n---\nbody"]
)
def test_whole_file_rejects_malformed_duplicate_and_alias_expansion(content):
    with pytest.raises(ValueError):
        parse_markdown(content, "source", 1, whole_file=True)


def test_callout_alias_block_anchor_and_complete_nested_edges():
    deps = {
        "Child": ApprovedDependency("child", 3, "content", "# H\n![[Grand#^b|Shown]]"),
        "Grand": ApprovedDependency("grand", 8, "content", "Text from grand. ^b"),
    }
    parsed = parse_markdown(
        "> [!NOTE]\n> ![[Child#H|Alias]]", "parent", 2, dependencies=deps, whole_file=True
    )
    # A child heading and its transitive body both retain their own original identities.
    assert [(e.source_identity, e.source_version) for e in parsed.elements] == [
        ("parent", 2),
        ("child", 3),
        ("child", 3),
        ("grand", 8),
    ]
    grand = parsed.elements[-1]
    assert grand.text == "Text from grand. ^b"
    assert grand.locator["char_start"] == 0 and grand.locator["char_end"] == len(grand.text)
    assert grand.locator["dependency_edges"][0]["parent_callouts"][0]["type"] == "NOTE"
    assert parsed.elements[0].text == "> [!NOTE]\n> ![[Child#H|Alias]]"
    assert (
        "".join(e.text for e in parsed.elements if e.source_identity == "child")
        == deps["Child"].text_content
    )
    assert [(edge["parent"], edge["child"]) for edge in grand.locator["dependency_edges"]] == [
        ("parent", "child"),
        ("child", "grand"),
    ]
    assert grand.locator["dependency_edges"][1]["child_version"] == 8


def test_revision_changes_and_roles_cannot_share_a_chunk():
    deps = {
        "old": ApprovedDependency("child", 1, "content", "OLD"),
        "new": ApprovedDependency("child", 2, "content", "NEW"),
    }
    parsed = parse_markdown("![[old]]\n\n![[new]]", "parent", 4, dependencies=deps)
    chunks = StructureFirstChunker().chunk(parsed)
    assert [(c.content, c.source_version) for c in chunks] == [("OLD", 1), ("NEW", 2)]
    deps["new"] = ApprovedDependency("child", 3, "content", "CHANGED")
    updated = StructureFirstChunker().chunk(
        parse_markdown("![[new]]", "parent", 4, dependencies=deps)
    )
    assert (updated[0].content, updated[0].source_version) == ("CHANGED", 3)
    for text in ("![[absent]]", "> [!NOTE]\n> ![[absent]]"):
        with pytest.raises(PermissionError):
            parse_markdown(text, "parent", 1, dependencies=deps)
    cycle = {"x": ApprovedDependency("child", 1, "content", "![[x]]")}
    with pytest.raises(ValueError, match="Cyclic"):
        parse_markdown("![[x]]", "parent", 1, dependencies=cycle)


def test_cumulative_transclusion_expansion_rejects_before_large_output():
    deps = {"x": ApprovedDependency("child", 1, "content", "a" * 80000)}
    with pytest.raises(BufferError, match="transclusion expansion"):
        parse_markdown("\n\n".join(["![[x]]"] * 32), "parent", 1, dependencies=deps)


def test_inert_multiline_comments_and_tilde_fences_do_not_resolve_children():
    text = "%% comment\n\n![[unapproved]]\n%%\n\n~~~\n![[also-inert]]\n~~~\n"
    parsed = parse_markdown(text, "source", 1, whole_file=True)
    content = "\n".join(e.text for e in parsed.elements)
    assert "![[unapproved]]" in content
    assert "![[also-inert]]" in content


def test_nested_heading_anchor_and_conflicting_same_revision_payload():
    deps = {"x": ApprovedDependency("child", 2, "content", "# Main\n## Sub\nSelected")}
    parsed = parse_markdown("![[x#main#sub|display]]", "root", 1, dependencies=deps)
    assert [e.text for e in parsed.elements] == ["Sub", "Selected"]
    deps["alias"] = ApprovedDependency("child", 2, "content", "Wrong same revision")
    with pytest.raises(ValueError, match="Conflicting child revision"):
        parse_markdown("![[x]]", "root", 1, dependencies=deps)


def test_docx_hyperlinks_nested_order_and_real_comment_ranges():
    source = docx.Document()
    p = source.add_paragraph("BEFORE ")
    link = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "VISIBLE_HYPERLINK"
    run.append(text)
    link.append(run)
    p._p.append(link)
    p.add_run(" AFTER")
    annotated = source.add_paragraph("Anchored text")
    source.add_comment(annotated.runs[0], "Real comment", author="Reader")
    outer = source.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.text = "CELL_BEFORE"
    inner = cell.add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "NESTED_MATERIAL"
    cell.paragraphs[-1].text = "CELL_AFTER"
    parsed = parse_docx(saved(source), "word", 7)
    body = [
        e
        for e in parsed.elements
        if e.kind != ElementKind.COMMENT and e.kind != ElementKind.METADATA
    ]
    assert [e.text for e in body] == [
        "BEFORE VISIBLE_HYPERLINK AFTER",
        "Anchored text",
        "CELL_BEFORE",
        "NESTED_MATERIAL",
        "CELL_AFTER",
    ]
    nested = next(e for e in body if e.text == "NESTED_MATERIAL")
    assert nested.locator["structural_path"] == [
        "body",
        "table",
        0,
        "row",
        0,
        "column",
        0,
        "table",
        0,
        "row",
        0,
        "column",
        0,
        "paragraph",
        0,
    ]
    comment = next(e for e in parsed.elements if e.kind == ElementKind.COMMENT)
    assert comment.locator["anchor"]["start"]["paragraph_index"] == 1
    assert comment.locator["anchor"]["start"]["char_offset"] == 0
    assert comment.locator["anchor"]["end"]["char_offset"] == len("Anchored text")
    assert all("page" not in e.locator for e in parsed.elements)


@pytest.mark.parametrize("identity", [None, "", " ", 0, []])
def test_docx_structural_identity_must_be_nonempty_typed(identity):
    e = DocumentElement(
        ElementKind.COMMENT,
        "x",
        "source",
        1,
        "content",
        {
            "source_key": "source",
            "source_version": 1,
            "approved_field": "content",
            "container_part": "comments",
            "comment_id": identity,
        },
    )
    with pytest.raises(ValueError):
        e.validate_mandatory_locator("DOCX")


def test_xlsx_formula_cache_merge_and_column_group_metadata():
    source = openpyxl.Workbook()
    sheet = source.active
    sheet.title = "Data"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Merged header"
    sheet["A2"] = "=SUM(1,2)"
    sheet["C2"] = "Grouped hidden"
    sheet.column_dimensions.group("B", "D", hidden=True)
    sheet.row_dimensions[2].hidden = True
    parsed = parse_xlsx(saved(source), "book", 2)
    header, row = parsed.elements
    assert header.locator["a1_range"] == "Data!A1:C1"
    assert [c["covered_cell"] for c in header.locator["cells"]] == [False, True, True]
    assert all(c["merged_range"] == "A1:C1" for c in header.locator["cells"])
    formula = row.locator["cells"][0]
    assert formula["formula"] == "=SUM(1,2)"
    assert formula["cache_state"] == "uncomputed"
    assert "=SUM(1,2)" in row.text
    assert row.locator["cells"][1]["coordinate"] == "C2"
    assert row.locator["cells"][1]["column_hidden"] is True
    assert row.locator["cells"][1]["row_hidden"] is True
    # A genuine cached value is preserved, without evaluating the formula.
    raw = saved(source)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as zin, zipfile.ZipFile(output, "w") as zout:
        for item in zin.infolist():
            content = zin.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"<f>SUM(1,2)</f><v></v>", b"<f>SUM(1,2)</f><v>3</v>")
            zout.writestr(item, content)
    cached = parse_xlsx(output.getvalue(), "book", 2).elements[1].locator["cells"][0]
    assert (cached["cache_state"], cached["cached_value"]) == ("cached", 3)


def test_xls_merges_and_unsupported_material_formula_fail_explicitly():
    source = xlwt.Workbook()
    sheet = source.add_sheet("Legacy")
    sheet.write_merge(0, 0, 0, 2, "Merged")
    sheet.col(1).hidden = True
    parsed = parse_xls(saved(source), "legacy", 3)
    cells = parsed.elements[0].locator["cells"]
    assert cells[0]["merged_range"] == "A1:C1"
    assert cells[1]["covered_cell"] is True
    assert cells[1]["column_hidden"] is True
    sheet.write(1, 0, xlwt.Formula("SUM(1,2)"))
    with pytest.raises(ValueError, match="UNSUPPORTED_MATERIAL_XLS_FORMULA"):
        parse_xls(saved(source), "legacy", 3)


def test_legacy_comment_record_cannot_be_silently_discarded():
    import struct

    from xlwt.CompoundDoc import XlsDoc

    source = xlwt.Workbook()
    source.add_sheet("S").write(0, 0, "Content")
    stream = source.get_biff_data()
    # Original NOTE identity with no verified companion drawing/text object.
    # The complete affected file must fail before xlrd can silently omit the note.
    note = struct.pack("<4H", 0, 0, 0, 1) + struct.pack("<HB", 6, 0) + b"Author\0"
    record = struct.pack("<HH", 0x001C, len(note)) + note
    output = io.BytesIO()
    XlsDoc().save(output, stream[:-4] + record + stream[-4:])
    with pytest.raises(ValueError, match="UNSUPPORTED_MATERIAL_XLS_COMMENT"):
        parse_xls(output.getvalue(), "legacy-comment", 1)


def test_real_scanned_merged_table_retains_measured_span():
    image = Image.new("RGB", (900, 480), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 32)
    for y in (60, 180, 300):
        draw.line((40, y, 840, y), fill="black", width=3)
    for x in (40, 840):
        draw.line((x, 60, x, 300), fill="black", width=3)
    draw.line((440, 180, 440, 300), fill="black", width=3)
    for text, pos in [("Item count", (65, 90)), ("Books", (65, 210)), ("45", (465, 210))]:
        draw.text(pos, text, font=font, fill="black")
    out = io.BytesIO()
    image.save(out, "PDF", resolution=144)
    parsed = parse_pdf(out.getvalue(), "merged-scan", 3)
    assert [e.text for e in parsed.elements] == ["Item count", "Books", "45"]
    assert parsed.elements[0].metadata["column_span"] == 2
    assert parsed.elements[0].locator["bbox"] == pytest.approx([20, 30, 420, 90], abs=2)


def test_archive_and_sparse_dimensions_reject_before_library_allocation(monkeypatch):
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"x" * 1_000_000)
    with pytest.raises(BufferError, match="archive member"):
        checked_zip(bomb.getvalue())
    source = openpyxl.Workbook()
    source.active["A1"] = "x"
    source.active["XFD10000"] = "x"
    data = saved(source)

    def forbidden(*args, **kwargs):
        raise AssertionError("Workbook allocated before admission")

    monkeypatch.setattr(openpyxl, "load_workbook", forbidden)
    with pytest.raises(BufferError, match="worksheet dimensions"):
        parse_xlsx(data, "sparse", 1)


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_package_entities_reject_before_document_library(encoding):
    xml = '<?xml version="1.0"?><!DOCTYPE root [<!ENTITY x "content">]><root>&x;</root>'
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("word/document.xml", xml.encode(encoding))
    with pytest.raises(ValueError, match="XML declaration"):
        parse_docx(output.getvalue(), "entity", 1)


def raster_pdf(table=False):
    image = Image.new("RGB", (900, 480), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 32)
    if table:
        for y in (60, 180, 300):
            draw.line((40, y, 840, y), fill="black", width=3)
        for x in (40, 440, 840):
            draw.line((x, 60, x, 300), fill="black", width=3)
        for text, pos in [
            ("Product", (65, 90)),
            ("Count", (465, 90)),
            ("Books", (65, 210)),
            ("45", (465, 210)),
        ]:
            draw.text(pos, text, font=font, fill="black")
    else:
        draw.text((60, 120), "Shelf contains 45 books.", font=font, fill="black")
    out = io.BytesIO()
    image.save(out, "PDF", resolution=144)
    return out.getvalue()


def test_real_scanned_table_has_original_geometry_and_cell_order():
    parsed = parse_pdf(raster_pdf(table=True), "scan", 4)
    assert [e.text for e in parsed.elements] == ["Product", "Count", "Books", "45"]
    assert parsed.table_count == 1
    assert [(e.locator["row_index"], e.locator["column_index"]) for e in parsed.elements] == [
        (0, 0),
        (0, 1),
        (1, 0),
        (1, 1),
    ]
    assert parsed.elements[0].locator["bbox"] == pytest.approx([20, 30, 220, 90], abs=2)
    assert all(e.locator["page"] == 1 and e.kind == ElementKind.TABLE_CELL for e in parsed.elements)


def test_native_pdf_paragraph_precedes_table():
    from test_pdf_geometry import pdf_with_stream

    commands = (
        b"BT /F1 12 Tf 50 750 Td (ABOVE_TABLE) Tj ET\n"
        b"50 600 m 250 600 l S 50 650 m 250 650 l S 50 700 m 250 700 l S\n"
        b"50 600 m 50 700 l S 150 600 m 150 700 l S 250 600 m 250 700 l S\n"
        b"BT /F1 12 Tf 60 675 Td (A) Tj ET\n"
        b"BT /F1 12 Tf 160 675 Td (B) Tj ET\n"
        b"BT /F1 12 Tf 60 625 Td (1) Tj ET\n"
        b"BT /F1 12 Tf 160 625 Td (2) Tj ET"
    )
    parsed = parse_pdf(pdf_with_stream(commands), "native", 2)
    assert [e.text for e in parsed.elements] == ["ABOVE_TABLE", "A", "B", "1", "2"]


@pytest.mark.parametrize(
    "vector",
    [np.zeros(384), np.ones(383), np.ones((1, 384)), np.full(384, np.nan), np.full(384, np.inf)],
)
def test_embedding_rejects_invalid_actual_output(monkeypatch, vector):
    pipeline = MultilingualEmbeddingPipeline(force_cpu=True)
    monkeypatch.setattr(pipeline, "_embed_onnx", lambda text: vector)
    with pytest.raises(ValueError):
        pipeline.embed_text("Invalid vector path")


def test_parser_to_embedding_on_representative_real_files_offline(monkeypatch):
    import socket

    def deny(*args, **kwargs):
        raise AssertionError("Unexpected Python network call")

    monkeypatch.setattr(socket, "socket", deny)
    model = MultilingualEmbeddingPipeline(force_cpu=True)
    word = docx.Document()
    word.add_paragraph("Nội dung tiếng Việt; 繁體中文; 简体中文.")
    excel = openpyxl.Workbook()
    excel.active["A1"] = "English spreadsheet"
    legacy = xlwt.Workbook()
    legacy.add_sheet("S").write(0, 0, "English legacy")
    documents = [
        parse_text("English TXT", "txt", 1),
        parse_markdown(WHOLE_FILE, "md", 1, whole_file=True),
        parse_docx(saved(word), "docx", 1),
        parse_xlsx(saved(excel), "xlsx", 1),
        parse_xls(saved(legacy), "xls", 1),
        parse_pdf(raster_pdf(), "pdf", 1),
    ]
    for document in documents:
        chunks = StructureFirstChunker(token_count=model.token_count).chunk(document)
        assert chunks
        vectors = model.embed_batch([c.content for c in chunks])
        assert len(vectors) == len(chunks)
        assert all(
            len(v) == 384 and math.isclose(sum(x * x for x in v), 1, abs_tol=1e-5) for v in vectors
        )


def test_real_parse_cooperative_cancellation_and_deadline():
    event = threading.Event()
    count = 0

    def cancelled():
        nonlocal count
        count += 1
        if count == 30:
            event.set()
        return event.is_set()

    with pytest.raises(ParsingCancelled), parsing_scope(cancelled):
        parse_text("line\n" * 100000, "large", 1)
    assert count == 30
    with pytest.raises(BufferError, match="deadline"), parsing_scope(seconds=0.001):
        time.sleep(0.003)
        parse_text("actual parser input", "source", 1)


def test_actual_worker_whole_file_and_malformed_encoding(monkeypatch):
    import ppv_rag.worker as worker
    from ppv_rag.config import Settings

    monkeypatch.setattr(
        worker,
        "settings",
        Settings(
            service_key="synthetic-test-service-credential-00000000",
            worker_dsn="postgresql://rag_worker@127.0.0.1:5432/vault",
        ),
    )
    monkeypatch.setattr(worker, "verify_document_grant", lambda request: None)
    staged = []
    monkeypatch.setattr(worker, "stage_document_result", lambda **kwargs: staged.append(kwargs))
    roots = [
        DocumentRootTask(
            root_key="whole",
            source_id="s1",
            format="md",
            content_text=WHOLE_FILE,
            approved_fields=["content"],
        )
    ]
    for fmt in ("md", "txt"):
        roots.append(
            DocumentRootTask(
                root_key=fmt,
                source_id=fmt,
                format=fmt,
                content_b64=base64.b64encode(b"bad\xfftext").decode(),
                approved_fields=["content"],
            )
        )
    request = DocumentBatchStageRequest(
        job_id="11111111-1111-1111-1111-111111111111",
        attempt_id="22222222-2222-2222-2222-222222222222",
        lease_token="a" * 64,
        roots=roots,
        input_digest=document_input_digest(roots),
    )
    response = worker.process_document_batch_request(request)
    assert (response.staged_count, response.failed_count) == (1, 2)
    content = "\n".join(c["content"] for c in staged[0]["chunks_payload"])
    assert all(
        s in content
        for s in ("private: true", "custom: value", "Whole comment", "<!-- Literal -->")
    )
    assert all(
        not r["chunks_payload"] and r["reason_code"] == "PARSER_VALIDATION_ERROR"
        for r in staged[1:]
    )


def test_worker_domain_projection_and_cancellation_during_actual_parser(
    bound_worker, claimed_unit_input, monkeypatch
):
    from concurrent.futures import ThreadPoolExecutor

    from conftest import HEADERS
    from fastapi.testclient import TestClient

    import ppv_rag.main as runtime
    import ppv_rag.parsers.text as text_parser
    import ppv_rag.worker as worker
    from ppv_rag.config import Settings

    monkeypatch.setattr(
        worker,
        "settings",
        Settings(
            service_key=bound_worker.service_key,
            worker_dsn="postgresql://rag_worker@127.0.0.1:5432/vault",
        ),
    )
    monkeypatch.setattr(worker, "verify_document_grant", lambda request: None)
    staged = []
    monkeypatch.setattr(worker, "stage_document_result", lambda **kwargs: staged.append(kwargs))
    roots = [
        DocumentRootTask(
            root_key="domain",
            source_id="domain",
            format="DOMAIN",
            prepared_fields=[
                {
                    "source_key": "domain",
                    "source_version": 1,
                    "approved_field": "title",
                    "text": "Java-approved title",
                }
            ],
        ),
        DocumentRootTask(
            root_key="injected",
            source_id="injected",
            format="DOMAIN",
            content_text="Unprepared text",
        ),
    ]
    request = DocumentBatchStageRequest(
        job_id=bound_worker.job_id,
        attempt_id=bound_worker.attempt_id,
        lease_token="a" * 64,
        roots=roots,
        input_digest=document_input_digest(roots),
    )
    result = worker.process_document_batch_request(request)
    assert (result.staged_count, result.failed_count) == (1, 1)
    assert staged[0]["chunks_payload"][0]["content"] == "Java-approved title"
    assert staged[1]["reason_code"] == "FIELD_NOT_APPROVED"
    staged.clear()

    entered, release = threading.Event(), threading.Event()
    original = text_parser.checkpoint
    calls = 0

    def checkpoint():
        nonlocal calls
        calls += 1
        if calls == 15:
            entered.set()
            assert release.wait(10)
        original()

    monkeypatch.setattr(text_parser, "checkpoint", checkpoint)
    roots = [
        DocumentRootTask(
            root_key="large",
            source_id="large",
            format="txt",
            content_text="Actual parsing line\n" * 10000,
            approved_fields=["content"],
        )
    ]
    request = request.model_copy(
        update={"roots": roots, "input_digest": document_input_digest(roots)}
    )
    with TestClient(runtime.app) as client, ThreadPoolExecutor(max_workers=1) as pool:
        future = pool.submit(
            client.post, "/v1/worker/execute-batch", json=request.model_dump(), headers=HEADERS
        )
        try:
            assert entered.wait(10)
            cancel = client.post(
                "/v1/worker/cancel", json={"job_id": request.job_id}, headers=HEADERS
            )
            assert cancel.status_code == 200
        finally:
            release.set()
        assert future.result(timeout=10).json()["status"] == "CANCELLED"
    assert calls == 15
    assert staged == []


def test_mixed_pdf_combines_native_and_real_scanned_cells():
    import pypdf
    from test_pdf_geometry import pdf_with_stream

    scanned = pypdf.PdfReader(io.BytesIO(raster_pdf(table=True)))
    native = pypdf.PdfReader(
        io.BytesIO(pdf_with_stream(b"BT /F1 12 Tf 15 225 Td (NATIVE_HEADING) Tj ET"))
    )
    writer = pypdf.PdfWriter()
    page = writer.add_page(scanned.pages[0])
    page.merge_page(native.pages[0])
    out = io.BytesIO()
    writer.write(out)
    parsed = parse_pdf(out.getvalue(), "mixed", 3)
    assert [e.text for e in parsed.elements] == [
        "NATIVE_HEADING",
        "Product",
        "Count",
        "Books",
        "45",
    ]
    assert parsed.elements[0].locator["is_ocr"] is False
    assert all(e.locator["is_ocr"] for e in parsed.elements[1:])


def test_overlapping_native_and_ocr_are_not_duplicated():
    import pypdf
    from test_pdf_geometry import pdf_with_stream

    image = Image.new("RGB", (900, 480), "white")
    out = io.BytesIO()
    image.save(out, "PDF", resolution=144)
    page = pypdf.PdfReader(io.BytesIO(out.getvalue())).pages[0]
    native = pypdf.PdfReader(io.BytesIO(pdf_with_stream(b"BT /F1 12 Tf 60 120 Td (OVERLAP) Tj ET")))
    writer = pypdf.PdfWriter()
    page = writer.add_page(page)
    page.merge_page(native.pages[0])
    out = io.BytesIO()
    writer.write(out)

    class PositionedOCR:
        def recognize_region(self, image_bytes, **kwargs):
            # The real selection path masks the original native region before recognition.
            image = Image.open(io.BytesIO(image_bytes))
            assert image.getpixel((130, 230))[:3] == (255, 255, 255)
            return OCRResult(blocks=[OCRTextBlock("OVERLAP", [60, 110, 115, 123], 0.99, "en")])

    parsed = parse_pdf(out.getvalue(), "overlap", 2, ocr_engine=PositionedOCR())
    assert [e.text for e in parsed.elements] == ["OVERLAP"]
    assert parsed.elements[0].locator["is_ocr"] is False


def test_pdf_inflation_rejects_before_pdf_library(monkeypatch):
    import zlib

    import ppv_rag.parsers.pdf as parser

    compressed = zlib.compress(b"x" * 9_000_000)
    data = (
        b"%PDF-1.4\n1 0 obj\n<< /Filter /FlateDecode >>\nstream\n"
        + compressed
        + b"\nendstream\nendobj\n%%EOF"
    )

    def forbidden(*args, **kwargs):
        raise AssertionError("PDF opened before inflation admission")

    monkeypatch.setattr(parser.pdfplumber, "open", forbidden)
    with pytest.raises(BufferError, match="PDF stream inflation"):
        parse_pdf(data, "bomb", 1)


@pytest.mark.parametrize(
    "dimensions", [b"/Width 50000 /Height 50000", b"/Height 50000 /Width 50000"]
)
def test_oversized_pdf_image_rejects_before_decoder_in_any_key_order(monkeypatch, dimensions):
    import ppv_rag.parsers.pdf as parser

    data = (
        b"%PDF-1.4\n1 0 obj\n<< /Subtype /Image /Filter /DCTDecode "
        + dimensions
        + b" >>\nstream\nencoded image\nendstream\nendobj\n%%EOF"
    )

    def forbidden(*args, **kwargs):
        raise AssertionError("PDF decoder allocated before dimension admission")

    monkeypatch.setattr(parser.pdfplumber, "open", forbidden)
    with pytest.raises(BufferError, match="PDF embedded raster"):
        parse_pdf(data, "oversized-image", 1)
