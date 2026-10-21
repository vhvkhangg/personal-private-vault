"""Comprehensive tests for TXT, Markdown/Obsidian, DOCX, XLSX, XLS, and PDF parsers."""

import io

import docx
import openpyxl
import pypdf
import pytest
import xlwt

from ppv_rag.parsers.docx import parse_docx
from ppv_rag.parsers.excel import parse_xls, parse_xlsx
from ppv_rag.parsers.markdown import ApprovedDependency, parse_markdown
from ppv_rag.parsers.models import ElementKind
from ppv_rag.parsers.pdf import parse_pdf
from ppv_rag.parsers.text import parse_text


def test_txt_parser_char_and_line_spans() -> None:
    content = "First paragraph line 1\nline 2\n\nSecond paragraph line 4"
    doc = parse_text(content, root_key="test-txt-1", source_version=1)
    assert doc.format_type == "TXT"
    assert len(doc.elements) == 2

    e1 = doc.elements[0]
    assert e1.kind == ElementKind.PARAGRAPH
    assert "First paragraph line 1\nline 2" in e1.text
    assert e1.locator["line_start"] == 1
    assert e1.locator["line_end"] == 2
    assert e1.locator["char_start"] == 0

    e2 = doc.elements[1]
    assert e2.kind == ElementKind.PARAGRAPH
    assert "Second paragraph line 4" in e2.text
    assert e2.locator["line_start"] == 4
    assert e2.locator["line_end"] == 4


def test_markdown_frontmatter_allowlist_and_comments() -> None:
    md = """---
title: "Approved Note"
tags: "personal, research"
unapproved_secret: "hidden_value"
---
%% Internal Obsidian comment %%
# Section 1
This is visible paragraph. <!-- HTML comment -->
"""
    doc = parse_markdown(
        md,
        root_key="test-md-1",
        source_version=1,
        approved_metadata_fields=frozenset({"title", "tags"}),
    )
    # Frontmatter element should only contain allowlisted fields
    assert any("Approved Note" in e.text and e.approved_field == "title" for e in doc.elements)
    assert not any("hidden_value" in e.text for e in doc.elements)
    assert any("unapproved_secret" in w for w in doc.warnings)

    # Heading element
    h1 = next(e for e in doc.elements if e.kind == ElementKind.HEADING)
    assert h1.kind == ElementKind.HEADING
    assert h1.text == "Section 1"
    assert h1.heading_path == ["Section 1"]

    # Paragraph element - comments stripped
    p = doc.elements[-1]
    assert p.text == "This is visible paragraph."
    assert "Internal Obsidian comment" not in p.text
    assert "HTML comment" not in p.text


def test_markdown_callout_and_block_ref() -> None:
    md = """# Guide
> [!NOTE]
> Important instruction for indexing.

Paragraph with block reference. ^block-123
"""
    doc = parse_markdown(md, root_key="test-md-2", source_version=1)
    callout = [e for e in doc.elements if e.kind == ElementKind.CALLOUT][0]
    assert "Important instruction" in callout.text
    assert callout.metadata["callout_type"] == "NOTE"

    para = [e for e in doc.elements if e.kind == ElementKind.PARAGRAPH][0]
    assert para.text == "Paragraph with block reference."
    assert para.locator["block_id"] == "block-123"


def test_markdown_transclusion_approved_vs_unauthorized() -> None:
    md = """# Parent Note
Here is an embed:
![[child-note-1]]
"""
    deps = {
        "child-note-1": ApprovedDependency(
            source_key="child-note-1",
            source_version=2,
            approved_field="content",
            text_content="Approved transcluded child content.",
        )
    }
    # Approved transclusion succeeds with child attribution
    doc = parse_markdown(md, root_key="parent-note", source_version=1, dependencies=deps)
    trans = [e for e in doc.elements if e.kind == ElementKind.TRANSCLUSION][0]
    assert trans.text == "Approved transcluded child content."
    assert trans.source_identity == "child-note-1"
    assert trans.source_version == 2
    assert trans.locator["parent_source_key"] == "parent-note"

    # Unapproved transclusion must raise PermissionError (reject file)
    md_unauthorized = """# Bad Note
![[unapproved-secret-note]]
"""
    with pytest.raises(PermissionError) as exc_info:
        parse_markdown(md_unauthorized, root_key="bad-note", source_version=1, dependencies=deps)
    assert "Unauthorized or unapproved transclusion" in str(exc_info.value)


def test_transitive_children_keep_exact_versions_and_reject_cycles() -> None:
    deps = {
        "child": ApprovedDependency("child", 4, "content", "![[grandchild]]"),
        "grandchild": ApprovedDependency("grandchild", 7, "content", "Visible grandchild."),
    }
    parsed = parse_markdown("![[child]]", "parent", 9, dependencies=deps)
    assert [(e.source_identity, e.source_version, e.text) for e in parsed.elements] == [
        ("grandchild", 7, "Visible grandchild.")
    ]
    cyclic = dict(deps, grandchild=ApprovedDependency("grandchild", 7, "content", "![[child]]"))
    with pytest.raises(ValueError, match="Cyclic"):
        parse_markdown("![[child]]", "parent", 9, dependencies=cyclic)


def test_unapproved_frontmatter_is_never_included_and_child_privacy_is_recursive() -> None:
    parsed = parse_markdown("---\ntitle: secret title\n---\nVisible body", "root", 1)
    assert all("secret title" not in e.text for e in parsed.elements)
    deps = {
        "child": ApprovedDependency("child", 3, "content", "![[grandchild]]"),
        "grandchild": ApprovedDependency(
            "grandchild", 8, "content", "---\nprivate: true\n---\nSecret"
        ),
    }
    with pytest.raises(ValueError, match="Restrictive"):
        parse_markdown("![[child]]", "parent", 1, dependencies=deps)


def test_txt_trimmed_offsets_still_address_exact_original_bytes() -> None:
    content = "  first paragraph  \n\n  second paragraph \n"
    parsed = parse_text(content, "source", 5)
    for element in parsed.elements:
        assert content[element.locator["char_start"] : element.locator["char_end"]] == element.text


def test_docx_parser() -> None:
    docx_file = docx.Document()
    docx_file.add_heading("Docx Title", level=1)
    docx_file.add_paragraph("Docx regular paragraph.")
    table = docx_file.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Col A"
    table.cell(0, 1).text = "Col B"
    table.cell(1, 0).text = "Val 1"
    table.cell(1, 1).text = "Val 2"

    bio = io.BytesIO()
    docx_file.save(bio)
    bio.seek(0)

    doc = parse_docx(bio.getvalue(), root_key="test-docx-1", source_version=1)
    assert doc.format_type == "DOCX"
    assert any(e.kind == ElementKind.HEADING and e.text == "Docx Title" for e in doc.elements)
    assert any(
        e.kind == ElementKind.PARAGRAPH and "regular paragraph" in e.text for e in doc.elements
    )
    assert any(e.kind == ElementKind.TABLE and "Col A" in e.text for e in doc.elements)
    assert doc.table_count == 1


def test_xlsx_parser() -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sales"
    ws1.append(["Product", "Revenue"])
    ws1.append(["Widget", 100])

    ws2 = wb.create_sheet("HiddenSheet")
    ws2.sheet_state = "hidden"
    ws2.append(["Secret", "Hidden"])

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    # Default include_hidden=True (ADR-0022 whole-file consent) extracts hidden sheet with metadata
    doc = parse_xlsx(bio.getvalue(), root_key="test-xlsx-1", source_version=1)
    assert doc.format_type == "XLSX"
    assert any("Product" in e.text and "Revenue" in e.text for e in doc.elements)
    assert any("Widget" in e.text and "100" in e.text for e in doc.elements)
    assert any("Secret" in e.text for e in doc.elements)
    assert any(e.metadata.get("sheet_hidden") is True for e in doc.elements)

    # With include_hidden=False, hidden sheet is excluded
    doc_no_hidden = parse_xlsx(
        bio.getvalue(), root_key="test-xlsx-1", source_version=1, include_hidden=False
    )
    assert not any("Secret" in e.text for e in doc_no_hidden.elements)
    assert any("HiddenSheet" in w for w in doc_no_hidden.warnings)

    # Mandatory locator check
    for elem in doc.elements:
        assert "worksheet" in elem.locator
        assert "a1_range" in elem.locator


def test_xlsx_rejects_material_row_truncation() -> None:
    wb = openpyxl.Workbook()
    wb.active.cell(10001, 1, "Must not silently omit this row")
    bio = io.BytesIO()
    wb.save(bio)
    with pytest.raises(BufferError, match="RESOURCE_LIMIT_EXCEEDED"):
        parse_xlsx(bio.getvalue(), "oversized", 1)


def test_xls_legacy_parser() -> None:
    wb = xlwt.Workbook()
    ws = wb.add_sheet("LegacySheet")
    ws.write(0, 0, "Item")
    ws.write(0, 1, "Cost")
    ws.write(1, 0, "Book")
    ws.write(1, 1, "25")

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    doc = parse_xls(bio.getvalue(), root_key="test-xls-1", source_version=1)
    assert doc.format_type == "XLS"
    assert any("Item" in e.text and "Cost" in e.text for e in doc.elements)
    assert any("Book" in e.text and "25" in e.text for e in doc.elements)
    for elem in doc.elements:
        assert elem.locator["worksheet"] == "LegacySheet"
        assert "a1_range" in elem.locator


def test_pdf_parser_native_text() -> None:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    # Write a simple PDF with an empty page
    bio = io.BytesIO()
    writer.write(bio)
    bio.seek(0)

    doc = parse_pdf(bio.getvalue(), root_key="test-pdf-1", source_version=1)
    assert doc.format_type == "PDF"
    assert doc.page_count == 1


def test_docx_hidden_runs_properties_and_comments() -> None:
    doc_file = docx.Document()
    doc_file.core_properties.title = "Secret Document"
    doc_file.core_properties.author = "Inspector"
    doc_file.core_properties.comments = "Internal review only"

    p1 = doc_file.add_paragraph("Visible start, ")
    run_hidden = p1.add_run("hidden run text, ")
    run_hidden.font.hidden = True
    p1.add_run("visible end.")

    doc_file.part.comments.add_comment(text="Needs citation", author="Reviewer")

    bio = io.BytesIO()
    doc_file.save(bio)
    bio.seek(0)

    # Default include_hidden=True extracts hidden runs, properties, comments
    doc = parse_docx(bio.getvalue(), root_key="test-docx-ext", source_version=1)
    assert doc.format_type == "DOCX"

    # Core properties present as METADATA
    title_elem = next(
        (e for e in doc.elements if e.kind == ElementKind.METADATA and "Secret Document" in e.text),
        None,
    )
    assert title_elem is not None
    assert title_elem.locator["container_part"] == "core_properties"
    assert title_elem.locator["property_key"] == "title"

    # Comments present as COMMENT
    comm_elem = next(
        (e for e in doc.elements if e.kind == ElementKind.COMMENT and "Needs citation" in e.text),
        None,
    )
    assert comm_elem is not None
    assert comm_elem.locator["container_part"] == "comments"
    assert "comment_id" in comm_elem.locator

    # Paragraph with hidden run
    para_elem = next(
        (e for e in doc.elements if e.kind == ElementKind.PARAGRAPH and "Visible start" in e.text),
        None,
    )
    assert para_elem is not None
    assert "hidden run text" in para_elem.text
    assert para_elem.metadata.get("has_hidden_runs") is True

    # With include_hidden=False, hidden run text is omitted
    doc_no_hidden = parse_docx(
        bio.getvalue(), root_key="test-docx-ext", source_version=1, include_hidden=False
    )
    para_no_hidden = next(
        (
            e
            for e in doc_no_hidden.elements
            if e.kind == ElementKind.PARAGRAPH and "Visible start" in e.text
        ),
        None,
    )
    assert para_no_hidden is not None
    assert "hidden run text" not in para_no_hidden.text


def test_xlsx_sparse_row_header_mapping() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory"
    # Row 1: Headers in columns A, B, C, D
    ws.cell(1, 1, "SKU")
    ws.cell(1, 2, "Name")
    ws.cell(1, 3, "Category")
    ws.cell(1, 4, "Qty")

    # Row 2: Sparse row with only Category and Qty (columns 3 and 4)
    ws.cell(2, 3, "Electronics")
    ws.cell(2, 4, 15)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    doc = parse_xlsx(bio.getvalue(), root_key="test-sparse", source_version=1)
    row2_elem = next((e for e in doc.elements if e.locator.get("row_start") == 2), None)
    assert row2_elem is not None
    # Category (col 3) must be paired with Category, NOT SKU!
    assert "Category: Electronics" in row2_elem.text
    assert "SKU: Electronics" not in row2_elem.text
    assert "Qty: 15" in row2_elem.text


def test_xlsx_hidden_rows_columns_and_comments() -> None:
    from openpyxl.comments import Comment as ExcelComment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.cell(1, 1, "ColA")
    ws.cell(1, 2, "ColB")
    ws.cell(2, 1, "ValA")
    ws.cell(2, 2, "ValB")
    ws.cell(3, 1, "HiddenRowVal")

    # Hide column B
    ws.column_dimensions["B"].hidden = True
    # Hide row 3
    ws.row_dimensions[3].hidden = True
    # Add comment to cell A1
    ws["A1"].comment = ExcelComment("Review cell A1", "Editor")

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    # include_hidden=True extracts hidden col, hidden row, and comment
    doc = parse_xlsx(bio.getvalue(), root_key="test-hidden", source_version=1, include_hidden=True)
    assert any("ColB" in e.text for e in doc.elements)
    assert any("HiddenRowVal" in e.text for e in doc.elements)
    assert any("Comment: Review cell A1" in e.text for e in doc.elements)

    # include_hidden=False skips hidden col and hidden row
    doc_no = parse_xlsx(
        bio.getvalue(), root_key="test-hidden", source_version=1, include_hidden=False
    )
    assert not any("ColB" in e.text for e in doc_no.elements)
    assert not any("HiddenRowVal" in e.text for e in doc_no.elements)


def test_unverified_vietnamese_backend_rejects_instead_of_warning() -> None:
    from ppv_rag.ocr.engine import AdaptiveOCREngine

    engine = AdaptiveOCREngine(default_language="en")
    with pytest.raises(ValueError, match="OCR_BACKEND_SCRIPT_UNVERIFIED"):
        engine.recognize_region(b"unread image", language="vi")
