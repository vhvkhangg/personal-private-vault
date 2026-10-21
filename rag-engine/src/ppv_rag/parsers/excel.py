"""Inert spreadsheet extraction with original cells, merge ancestry and finite admission."""

from __future__ import annotations

import io
import struct
import zipfile
from collections import defaultdict
from typing import Any, BinaryIO
from xml.etree import ElementTree as ET

import openpyxl  # type: ignore[import-untyped]
import xlrd  # type: ignore[import-untyped]
from openpyxl.utils.cell import range_boundaries  # type: ignore[import-untyped]

from ppv_rag.limits import (
    MAX_CELLS,
    MAX_RECTANGULAR_CELLS,
    check_text,
    checked_zip,
    checkpoint,
    input_bytes,
)
from ppv_rag.parsers.models import DocumentElement, ElementKind, ParsedDocument


def _col_letter(col_idx: int) -> str:
    result = ""
    while col_idx > 0:
        col_idx, remainder = divmod(col_idx - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _preflight_xlsx(data: bytes) -> None:
    cells = expanded = 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist():
            if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                continue
            with archive.open(name) as stream:
                for _, node in ET.iterparse(stream, events=("end",)):
                    checkpoint()
                    tag = node.tag.rsplit("}", 1)[-1]
                    if tag in {"dimension", "mergeCell"}:
                        ref = node.attrib.get("ref", "")
                        c1, r1, c2, r2 = range_boundaries(ref)
                        if not all(isinstance(x, int) for x in (c1, r1, c2, r2)):
                            raise ValueError("Invalid worksheet bounds")
                        area = (c2 - c1 + 1) * (r2 - r1 + 1)
                        if r2 > 10000 or c2 > 16384 or area > MAX_RECTANGULAR_CELLS:
                            raise BufferError("RESOURCE_LIMIT_EXCEEDED: worksheet dimensions")
                        if tag == "mergeCell":
                            expanded += area
                    if tag == "c":
                        cells += 1
                        c1, r1, _, _ = range_boundaries(node.attrib["r"])
                        if r1 > 10000 or c1 > 16384:
                            raise BufferError("RESOURCE_LIMIT_EXCEEDED: cell coordinates")
                    if cells + expanded > MAX_CELLS:
                        raise BufferError("RESOURCE_LIMIT_EXCEEDED: worksheet cells")
                    node.clear()


def parse_xlsx(
    file_bytes: bytes | BinaryIO,
    root_key: str,
    source_version: int,
    approved_field: str = "content",
    include_hidden: bool = True,
) -> ParsedDocument:
    data = checked_zip(file_bytes)
    _preflight_xlsx(data)  # Before openpyxl allocates cells or expands merges.
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    doc = ParsedDocument(root_key, "XLSX", source_version)
    total_text = 0
    try:
        for sheet in wb:
            checkpoint()
            sheet_hidden = sheet.sheet_state != "visible"
            if sheet_hidden and not include_hidden:
                doc.warnings.append(f"Skipped hidden worksheet: {sheet.title}")
                continue
            groups = [
                (d.min, d.max, bool(d.hidden), d.outlineLevel)
                for d in sheet.column_dimensions.values()
            ]
            merged: dict[tuple[int, int], dict[str, Any]] = {}
            for region in sheet.merged_cells.ranges:
                for row in range(region.min_row, region.max_row + 1):
                    for col in range(region.min_col, region.max_col + 1):
                        merged[(row, col)] = {
                            "merged_range": str(region),
                            "merged_anchor": f"{_col_letter(region.min_col)}{region.min_row}",
                            "covered_cell": (row, col) != (region.min_row, region.min_col),
                        }
            rows: dict[int, list[Any]] = defaultdict(list)
            # Sparse iteration over loaded original cells, never rectangular iter_rows.
            for (row, _), cell in sorted(sheet._cells.items()):
                if (
                    cell.value is not None
                    or cell.comment is not None
                    or (row, cell.column) in merged
                ):
                    rows[row].append(cell)
            headers: dict[int, str] = {}
            for row_index, cells in sorted(rows.items()):
                checkpoint()
                row_dim = sheet.row_dimensions.get(row_index)
                row_hidden = bool(row_dim and row_dim.hidden)
                if row_hidden and not include_hidden:
                    continue
                values: list[tuple[int, str]] = []
                cell_records: list[dict[str, Any]] = []
                comments = {}
                for cell in cells:
                    matching = [g for g in groups if g[0] <= cell.column <= g[1]]
                    col_hidden = any(g[2] for g in matching)
                    if col_hidden and not include_hidden:
                        continue
                    value = "" if cell.value is None else str(cell.value)
                    record = {
                        "coordinate": cell.coordinate,
                        "value": value,
                        "sheet_hidden": sheet_hidden,
                        "row_hidden": row_hidden,
                        "column_hidden": col_hidden,
                        "column_groups": matching,
                        "row_outline_level": row_dim.outlineLevel if row_dim else 0,
                        **merged.get((row_index, cell.column), {}),
                    }
                    if cell.data_type == "f":
                        if not isinstance(cell.value, str):
                            raise ValueError("Unsupported material formula representation")
                        cache = cached[sheet.title][cell.coordinate].value
                        record.update(
                            formula=cell.value,
                            cached_value=cache,
                            cache_state="uncomputed" if cache is None else "cached",
                        )
                        value = f"{cell.value} [cache: {record['cache_state']}"
                        value += f"; value: {cache}]" if cache is not None else "]"
                    if cell.comment:
                        comments[cell.coordinate] = cell.comment.text
                        record["comment"] = cell.comment.text
                        value += f" [Comment: {cell.comment.text}]"
                    if value:
                        values.append((cell.column, value))
                    cell_records.append(record)
                if not values:
                    continue
                is_header = not headers
                if is_header:
                    headers = dict(values)
                    text = f"Worksheet [{sheet.title}] Header: " + " | ".join(v for _, v in values)
                else:
                    text = f"[{sheet.title} row {row_index}] " + " | ".join(
                        f"{headers.get(c, f'Col{_col_letter(c)}')}: {v}" for c, v in values
                    )
                cols = [
                    openpyxl.utils.cell.coordinate_from_string(c["coordinate"])[0]
                    for c in cell_records
                ]
                indices = [openpyxl.utils.cell.column_index_from_string(c) for c in cols]
                first, last = min(indices), max(indices)
                locator = {
                    "source_key": root_key,
                    "source_version": source_version,
                    "approved_field": approved_field,
                    "worksheet": sheet.title,
                    "a1_range": f"{sheet.title}!{_col_letter(first)}{row_index}:"
                    f"{_col_letter(last)}{row_index}",
                    "row_start": row_index,
                    "row_end": row_index,
                    "col_start": first,
                    "col_end": last,
                    "is_header": is_header,
                    "cells": cell_records,
                }
                meta = {
                    "sheet_hidden": sheet_hidden,
                    "row_hidden": row_hidden,
                    "is_hidden": any(c["column_hidden"] for c in cell_records)
                    or row_hidden
                    or sheet_hidden,
                    "comments": comments,
                    "cells": cell_records,
                }
                total_text += len(text)
                if total_text > 1_000_000:
                    raise BufferError("RESOURCE_LIMIT_EXCEEDED: workbook text")
                element = DocumentElement(
                    ElementKind.SPREADSHEET_ROW,
                    text,
                    root_key,
                    source_version,
                    approved_field,
                    locator,
                    metadata=meta,
                )
                element.validate_mandatory_locator("XLSX")
                doc.elements.append(element)
            doc.table_count += 1
        check_text("\n".join(e.text for e in doc.elements))
        return doc
    finally:
        wb.close()
        cached.close()


def _preflight_xls(data: bytes) -> None:
    from xlrd.compdoc import CompDoc  # type: ignore[import-untyped]

    # OLE streams are bounded by the admitted byte array; no archive extraction on disk.
    package = CompDoc(data, logfile=io.StringIO())
    stream = package.get_named_stream("Workbook") or package.get_named_stream("Book")
    if stream is None:
        raise ValueError("Missing workbook stream")
    offset = records = 0
    while offset + 4 <= len(stream):
        checkpoint()
        opcode, size = struct.unpack_from("<HH", stream, offset)
        offset += 4
        record = stream[offset : offset + size]
        if len(record) != size:
            raise ValueError("Truncated workbook record")
        if opcode == 0x0809 and size >= 2 and struct.unpack_from("<H", record)[0] != 0x0600:
            raise ValueError("UNSUPPORTED_MATERIAL_XLS_BIFF_VERSION")
        if opcode == 0x00FC and size >= 8:
            if struct.unpack_from("<I", record, 4)[0] > MAX_CELLS:
                raise BufferError("RESOURCE_LIMIT_EXCEEDED: legacy shared strings")
        if opcode in (0x0006, 0x0206, 0x0406):
            # xlrd exposes only cached values, not exact formula text. Reject instead of losing it.
            raise ValueError("UNSUPPORTED_MATERIAL_XLS_FORMULA")
        if opcode in (0x001C, 0x01B6):
            # xlrd drops NOTE records without a matching drawing/TXO object. Until
            # that complete graph is verified, cached/partial comments are unsafe.
            raise ValueError("UNSUPPORTED_MATERIAL_XLS_COMMENT")
        if opcode == 0x0200 and size >= 12:
            rows = struct.unpack_from("<I", record, 4)[0]
            cols = struct.unpack_from("<H", record, 10)[0]
            if rows > 10000 or rows * cols > MAX_RECTANGULAR_CELLS:
                raise BufferError("RESOURCE_LIMIT_EXCEEDED: legacy dimensions")
        records += 1
        if records > 200000:
            raise BufferError("RESOURCE_LIMIT_EXCEEDED: legacy records")
        offset += size


def parse_xls(
    file_bytes: bytes,
    root_key: str,
    source_version: int,
    approved_field: str = "content",
    include_hidden: bool = True,
) -> ParsedDocument:
    """Parse legacy XLS workbook via xlrd with formatting_info=True.

    Supports extracting hidden sheets, rows, and columns when include_hidden=True.
    Accurately associates sparse-row columns with header definitions by column index.
    """
    file_bytes = input_bytes(file_bytes)
    _preflight_xls(file_bytes)
    wb = xlrd.open_workbook(
        file_contents=file_bytes, formatting_info=True, ragged_rows=True, logfile=io.StringIO()
    )
    doc = ParsedDocument(root_key=root_key, format_type="XLS", source_version=source_version)
    total_text = 0

    for sheet_idx in range(wb.nsheets):
        sheet = wb.sheet_by_index(sheet_idx)
        # xlrd visibility: 0=visible, 1=hidden, 2=very hidden
        sheet_hidden = sheet.visibility != 0
        if sheet_hidden:
            if not include_hidden:
                doc.warnings.append(f"Skipped hidden worksheet: {sheet.name}")
                continue

        if sheet.nrows == 0:
            continue

        if sheet.nrows > 10000 or sheet.nrows * sheet.ncols > MAX_CELLS:
            raise BufferError("RESOURCE_LIMIT_EXCEEDED: Sheet row count exceeds 10,000")

        # Check hidden columns via colinfo_map
        hidden_cols: set[int] = set()
        colinfo_map = getattr(sheet, "colinfo_map", {})
        for col_idx, colinfo in colinfo_map.items():
            if getattr(colinfo, "hidden", 0) == 1:
                hidden_cols.add(col_idx)

        # Check hidden rows via rowinfo_map
        hidden_rows: set[int] = set()
        rowinfo_map = getattr(sheet, "rowinfo_map", {})
        for row_idx, rowinfo in rowinfo_map.items():
            if getattr(rowinfo, "hidden", 0) == 1:
                hidden_rows.add(row_idx)

        doc.table_count += 1
        headers_by_col: dict[int, str] = {}

        for r_idx in range(sheet.nrows):
            checkpoint()
            row_hidden = r_idx in hidden_rows
            if row_hidden and not include_hidden:
                continue
            one_based_row = r_idx + 1
            visible_cells: list[tuple[int, str, bool]] = []
            cell_records: list[dict[str, Any]] = []
            for c in range(sheet.row_len(r_idx)):
                col_hidden = c in hidden_cols
                if col_hidden and not include_hidden:
                    continue
                v_str = str(sheet.cell_value(r_idx, c)).strip()
                note = sheet.cell_note_map.get((r_idx, c))
                if note:
                    v_str += f" [Comment: {note.text}]"
                record: dict[str, Any] = {
                    "coordinate": f"{_col_letter(c + 1)}{r_idx + 1}",
                    "column_hidden": col_hidden,
                    "row_hidden": row_hidden,
                    "sheet_hidden": sheet_hidden,
                    "comment": note.text if note else None,
                }
                for rlo, rhi, clo, chi in sheet.merged_cells:
                    if rlo <= r_idx < rhi and clo <= c < chi:
                        record.update(
                            merged_range=f"{_col_letter(clo + 1)}{rlo + 1}:{_col_letter(chi)}{rhi}",
                            merged_anchor=f"{_col_letter(clo + 1)}{rlo + 1}",
                            covered_cell=(r_idx, c) != (rlo, clo),
                        )
                cell_records.append(record)
                if v_str:
                    is_cell_hidden = col_hidden or row_hidden or sheet_hidden
                    visible_cells.append((c + 1, v_str, is_cell_hidden))

            if not visible_cells:
                continue

            min_col = visible_cells[0][0]
            max_col = visible_cells[-1][0]
            start_col_letter = _col_letter(min_col)
            end_col_letter = _col_letter(max_col)
            range_coords = f"{start_col_letter}{one_based_row}:{end_col_letter}{one_based_row}"
            a1_range = f"{sheet.name}!{range_coords}"

            row_has_hidden = any(h for _, _, h in visible_cells)

            if not headers_by_col:
                for c_idx, val, _ in visible_cells:
                    headers_by_col[c_idx] = val
                header_joined = " | ".join(val for _, val, _ in visible_cells)
                text = f"Worksheet [{sheet.name}] Header: {header_joined}"
                meta: dict[str, Any] = {
                    "worksheet": sheet.name,
                    "row": one_based_row,
                    "is_header": True,
                }
                if sheet_hidden:
                    meta["sheet_hidden"] = True
                if row_hidden:
                    meta["row_hidden"] = True
                if row_has_hidden:
                    meta["is_hidden"] = True

                total_text += len(text)
                if total_text > 1_000_000:
                    raise BufferError("RESOURCE_LIMIT_EXCEEDED: workbook text")
                elem = DocumentElement(
                    kind=ElementKind.SPREADSHEET_ROW,
                    text=text,
                    source_identity=root_key,
                    source_version=source_version,
                    approved_field=approved_field,
                    locator={
                        "source_key": root_key,
                        "source_version": source_version,
                        "approved_field": approved_field,
                        "worksheet": sheet.name,
                        "a1_range": a1_range,
                        "cells": cell_records,
                        "row_start": one_based_row,
                        "row_end": one_based_row,
                        "col_start": min_col,
                        "col_end": max_col,
                        "is_header": True,
                    },
                    metadata=meta,
                )
                elem.validate_mandatory_locator("XLS")
                doc.elements.append(elem)
            else:
                paired = [
                    f"{headers_by_col.get(c_idx, f'Col{_col_letter(c_idx)}')}: {val}"
                    for c_idx, val, _ in visible_cells
                ]
                text = f"[{sheet.name} row {one_based_row}] " + " | ".join(paired)
                meta = {
                    "worksheet": sheet.name,
                    "row": one_based_row,
                    "is_header": False,
                }
                if sheet_hidden:
                    meta["sheet_hidden"] = True
                if row_hidden:
                    meta["row_hidden"] = True
                if row_has_hidden:
                    meta["is_hidden"] = True

                total_text += len(text)
                if total_text > 1_000_000:
                    raise BufferError("RESOURCE_LIMIT_EXCEEDED: workbook text")
                elem = DocumentElement(
                    kind=ElementKind.SPREADSHEET_ROW,
                    text=text,
                    source_identity=root_key,
                    source_version=source_version,
                    approved_field=approved_field,
                    locator={
                        "source_key": root_key,
                        "source_version": source_version,
                        "approved_field": approved_field,
                        "worksheet": sheet.name,
                        "a1_range": a1_range,
                        "cells": cell_records,
                        "row_start": one_based_row,
                        "row_end": one_based_row,
                        "col_start": min_col,
                        "col_end": max_col,
                        "is_header": False,
                    },
                    metadata=meta,
                )
                elem.validate_mandatory_locator("XLS")
                doc.elements.append(elem)

    wb.release_resources()
    check_text("\n".join(e.text for e in doc.elements))
    return doc
